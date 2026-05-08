"""
Bulk client import from CSV or Excel (.xlsx).

Handles encoding, delimiters, Excel cell types, validation, deduplication, and limits.
"""
from __future__ import annotations

import csv
import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from typing import BinaryIO, Iterable

from django.core.exceptions import ValidationError
from django.core.validators import EmailValidator
from django.db import transaction

from .models import Client

# --- Limits (abuse prevention + predictable memory) ---
MAX_FILE_BYTES = 5 * 1024 * 1024  # 5 MiB
MAX_DATA_ROWS = 5000
# Match model / reasonable caps
MAX_NAME_LEN = 255
MAX_EMAIL_LEN = 254
MAX_PHONE_LEN = 20
MAX_COMPANY_LEN = 255
MAX_ADDRESS_LEN = 10000

# Headers normalized (lowercase, internal whitespace collapsed) -> canonical field
HEADER_ALIASES: dict[str, str] = {
    "name": "name",
    "full name": "name",
    "client name": "name",
    "customer name": "name",
    "contact name": "name",
    "email": "email",
    "e-mail": "email",
    "email address": "email",
    "phone": "phone",
    "mobile": "phone",
    "telephone": "phone",
    "tel": "phone",
    "cell": "phone",
    "company": "company",
    "organization": "company",
    "organisation": "company",
    "org": "company",
    "business": "company",
    "address": "address",
}


def normalize_header_label(raw: str) -> str:
    """Strip BOM/whitespace, lowercase, collapse internal whitespace."""
    if raw is None:
        return ""
    text = str(raw).replace("\ufeff", "").strip()
    text = re.sub(r"\s+", " ", text)
    return text.lower()


def normalize_email_key(email: str) -> str:
    """Normalize for dedupe / DB lookup (empty -> '')."""
    e = email.strip().lower()
    return e


class ImportParseError(Exception):
    """Raised when the file cannot be read (format, encoding, structure)."""

    def __init__(self, message: str):
        self.message = message
        super().__init__(message)


@dataclass
class RowData:
    """One logical row after parsing (before validation)."""

    sheet_row: int
    raw: dict[str, str]


@dataclass
class ImportResult:
    created: int = 0
    skipped_duplicate_in_file: list[dict] = field(default_factory=list)
    skipped_existing_email: list[dict] = field(default_factory=list)
    row_errors: list[dict] = field(default_factory=list)
    # True when the file parsed but every data row was blank (or only a header was present)
    had_no_data_rows: bool = False


def _cell_to_str(value: object) -> str:
    """
    Convert Excel/CSV cell to a string suitable for Client CharField/TextField.

    Handles numbers (including Excel-stored phone numbers), dates, bool, None.
    """
    if value is None:
        return ""
    # Excel checkbox / boolean is uncommon for contact fields; avoid polluting text columns.
    if isinstance(value, bool):
        return ""
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if not math.isfinite(value):
            return ""
        # Whole numbers: avoid 1.23e+10 scientific notation in output
        if abs(value - round(value)) < 1e-9:
            try:
                iv = int(round(value))
                return str(iv)
            except (OverflowError, ValueError):
                pass
        # Prefer plain digits for large floats (Excel phone edge case)
        formatted = f"{value:.10f}".rstrip("0").rstrip(".")
        return formatted if formatted else str(value)
    if isinstance(value, Decimal):
        if value == value.to_integral_value():
            return str(int(value))
        return format(value, "f").rstrip("0").rstrip(".")
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    s = str(value)
    return s.replace("\x00", "").strip()


def _truncate_field(field_name: str, value: str, max_len: int) -> tuple[str, str | None]:
    """Returns (maybe_truncated, error_message_if_invalid)."""
    if len(value) > max_len:
        return value[:max_len], f"{field_name} exceeds maximum length ({max_len} characters)."
    return value, None


def _validate_row_fields(
    name: str,
    email: str,
    phone: str,
    company: str,
    address: str,
) -> list[str]:
    """Return a list of human-readable errors (empty if valid)."""
    errors: list[str] = []
    email_validator = EmailValidator()

    if not name.strip():
        errors.append("Name is required.")
    else:
        _, err = _truncate_field("Name", name.strip(), MAX_NAME_LEN)
        if err:
            errors.append(err)

    if email.strip():
        e = email.strip()
        if len(e) > MAX_EMAIL_LEN:
            errors.append(f"Email exceeds maximum length ({MAX_EMAIL_LEN} characters).")
        else:
            try:
                email_validator(e)
            except ValidationError:
                errors.append("Enter a valid email address.")

    if phone.strip():
        _, err = _truncate_field("Phone", phone.strip(), MAX_PHONE_LEN)
        if err:
            errors.append(err)

    if company.strip():
        _, err = _truncate_field("Company", company.strip(), MAX_COMPANY_LEN)
        if err:
            errors.append(err)

    if address.strip():
        _, err = _truncate_field("Address", address.strip(), MAX_ADDRESS_LEN)
        if err:
            errors.append(err)

    return errors


def _build_column_map(header_cells: list[str]) -> dict[str, int]:
    """
    Map canonical field name -> column index.
    Raises ImportParseError on missing name column or duplicate mappings.
    """
    canonical_to_index: dict[str, int] = {}

    for col_idx, raw_h in enumerate(header_cells):
        label = normalize_header_label(_cell_to_str(raw_h))
        if not label:
            continue
        canonical = HEADER_ALIASES.get(label)
        if not canonical:
            # Ignore unknown columns so exports with extra fields still import.
            continue
        if canonical in canonical_to_index:
            raise ImportParseError(
                f'Duplicate column for "{canonical}". Remove duplicate headers '
                f"(columns {canonical_to_index[canonical] + 1} and {col_idx + 1})."
            )
        canonical_to_index[canonical] = col_idx

    if "name" not in canonical_to_index:
        raise ImportParseError(
            'Missing required column "name". Use the provided template or include a '
            '"name" column (aliases: full name, client name, customer name).'
        )

    return canonical_to_index


def _row_is_empty(values: dict[str, str]) -> bool:
    return not any((v or "").strip() for v in values.values())


def _parse_csv(file: BinaryIO) -> list[RowData]:
    raw = file.read()
    if not raw:
        raise ImportParseError("The file is empty.")

    if len(raw) > MAX_FILE_BYTES:
        raise ImportParseError(
            f"File is too large. Maximum size is {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )

    # UTF-8 with BOM (Excel-exported CSV) and strict UTF-8
    for encoding in ("utf-8-sig", "utf-8"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            text = None
    else:
        raise ImportParseError(
            "Could not decode CSV as UTF-8. Save the file as UTF-8 and try again."
        )

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
    except csv.Error:
        dialect = csv.excel

    # Handle "\r\n" only files
    stream = io.StringIO(text)
    reader = csv.reader(stream, dialect)

    rows_iter = iter(reader)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        raise ImportParseError("The CSV has no header row.")

    column_map = _build_column_map(header_row)
    out: list[RowData] = []
    data_row_count = 0
    # Row 1 in file = header; first data row is spreadsheet row 2
    sheet_row = 2

    for parts in rows_iter:
        values: dict[str, str] = {}
        for canonical, col_idx in column_map.items():
            cell = parts[col_idx] if col_idx < len(parts) else ""
            values[canonical] = _cell_to_str(cell)

        if _row_is_empty(values):
            sheet_row += 1
            continue

        data_row_count += 1
        if data_row_count > MAX_DATA_ROWS:
            raise ImportParseError(
                f"Too many data rows. Maximum is {MAX_DATA_ROWS} rows per import."
            )

        out.append(RowData(sheet_row=sheet_row, raw=values))
        sheet_row += 1

    return out


def _parse_xlsx(file: BinaryIO) -> list[RowData]:
    raw = file.read()
    if not raw:
        raise ImportParseError("The file is empty.")

    if len(raw) > MAX_FILE_BYTES:
        raise ImportParseError(
            f"File is too large. Maximum size is {MAX_FILE_BYTES // (1024 * 1024)} MB."
        )

    if not zipfile.is_zipfile(io.BytesIO(raw)):
        raise ImportParseError("Invalid Excel file (.xlsx). Export as .xlsx or use CSV.")

    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise ImportParseError("Excel support is not installed on the server.") from exc

    bio = io.BytesIO(raw)
    try:
        wb = load_workbook(filename=bio, read_only=True, data_only=True)
    except Exception as exc:
        raise ImportParseError("Could not read Excel file. Use .xlsx format (Excel 2007+).") from exc

    try:
        ws = wb.active
        rows_iter: Iterable[tuple] = ws.iter_rows(values_only=True)
        try:
            header_row = next(rows_iter)
        except StopIteration:
            raise ImportParseError("The spreadsheet has no header row.")

        header_cells = list(header_row)
        column_map = _build_column_map([_cell_to_str(c) for c in header_cells])

        out: list[RowData] = []
        data_row_count = 0
        sheet_row = 2

        for row in rows_iter:
            parts = list(row)
            values = {}
            for canonical, col_idx in column_map.items():
                cell_val = parts[col_idx] if col_idx < len(parts) else None
                values[canonical] = _cell_to_str(cell_val)

            if _row_is_empty(values):
                sheet_row += 1
                continue

            data_row_count += 1
            if data_row_count > MAX_DATA_ROWS:
                raise ImportParseError(
                    f"Too many data rows. Maximum is {MAX_DATA_ROWS} rows per import."
                )

            out.append(RowData(sheet_row=sheet_row, raw=values))
            sheet_row += 1

        return out
    finally:
        wb.close()


def parse_client_import_file(file: BinaryIO, filename: str) -> list[RowData]:
    """
    Parse CSV or XLSX into row structs. Filename extension determines format.
    """
    name = (filename or "").lower().strip()
    if name.endswith(".csv"):
        return _parse_csv(file)
    if name.endswith(".xlsx"):
        return _parse_xlsx(file)
    if name.endswith(".xls"):
        raise ImportParseError(
            "Old Excel format (.xls) is not supported. Save as .xlsx or export as CSV."
        )
    raise ImportParseError(
        "Unsupported file type. Upload a .csv or .xlsx file."
    )


def import_clients_for_user(user, rows: list[RowData]) -> ImportResult:
    """
    Validate rows, dedupe by email, skip rows that match an existing client email, bulk insert.
    All-or-nothing is NOT used: valid rows are saved; invalid rows are reported.
    """
    result = ImportResult()

    # Existing emails for this user (case-insensitive), non-empty only
    existing_qs = Client.objects.filter(user=user).exclude(email="").values_list("email", flat=True)
    existing_normalized: set[str] = {normalize_email_key(e) for e in existing_qs if e}

    seen_in_file: set[str] = set()
    to_create: list[Client] = []

    for rd in rows:
        raw = rd.raw
        name = (raw.get("name") or "").strip()
        email_raw = (raw.get("email") or "").strip()
        phone = (raw.get("phone") or "").strip()
        company = (raw.get("company") or "").strip()
        address = (raw.get("address") or "").strip()

        field_errors = _validate_row_fields(name, email_raw, phone, company, address)
        if field_errors:
            for msg in field_errors:
                result.row_errors.append({"row": rd.sheet_row, "message": msg})
            continue

        email_norm = normalize_email_key(email_raw)
        if email_norm:
            if email_norm in seen_in_file:
                result.skipped_duplicate_in_file.append(
                    {
                        "row": rd.sheet_row,
                        "email": email_raw,
                        "message": "Duplicate email in this file; earlier row was kept.",
                    }
                )
                continue
            seen_in_file.add(email_norm)

            if email_norm in existing_normalized:
                result.skipped_existing_email.append(
                    {
                        "row": rd.sheet_row,
                        "email": email_raw,
                        "message": "A client with this email already exists; skipped.",
                    }
                )
                continue

        # Truncate after validation passed length checks — safety if constants drift
        name = name[:MAX_NAME_LEN]
        email = email_raw[:MAX_EMAIL_LEN] if email_raw else ""
        phone = phone[:MAX_PHONE_LEN]
        company = company[:MAX_COMPANY_LEN]
        address = address[:MAX_ADDRESS_LEN]

        to_create.append(
            Client(
                user=user,
                name=name,
                email=email,
                phone=phone,
                company=company,
                address=address,
            )
        )
        if email_norm:
            # Reserve email for rest of this batch so two new rows don't duplicate
            existing_normalized.add(email_norm)

    if not to_create:
        return result

    with transaction.atomic():
        Client.objects.bulk_create(to_create, batch_size=500)

    result.created = len(to_create)
    return result


def run_client_import(user, file: BinaryIO, filename: str) -> ImportResult:
    """Parse + import in one call."""
    rows = parse_client_import_file(file, filename)
    if not rows:
        return ImportResult(had_no_data_rows=True)
    return import_clients_for_user(user, rows)


def client_import_template_xlsx_bytes() -> bytes:
    """Sample .xlsx for download (UTF-8 not applicable; single sheet)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["name", "email", "phone", "company", "address"])
    ws.append(
        [
            "Example Client",
            "client@example.com",
            "+1 555-0100",
            "Example Co",
            "123 Main Street",
        ]
    )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def client_import_template_csv_bytes() -> bytes:
    """Sample CSV for download (UTF-8 with BOM for Excel compatibility)."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["name", "email", "phone", "company", "address"])
    writer.writerow(
        [
            "Example Client",
            "client@example.com",
            "+1 555-0100",
            "Example Co",
            "123 Main Street",
        ]
    )
    return buf.getvalue().encode("utf-8-sig")
